import os
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 1. 워드 문서 생성 및 스타일 설정
doc = Document()

# 문서 제목 설정
title = doc.add_heading('운동 처치에 따른 VO2max 및 혈중 젖산 농도 변화 연구 보고서', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('작성일: 2026년 | 연구 기관: 충북스포츠과학센터 | 연구자: youngsam032')
doc.add_paragraph().paragraph_format.space_after = Pt(12)

# 2. 요약 통계표 삽입
doc.add_heading('1. 주요 연구 변인 요약 통계', level=1)

if os.path.exists('raw_exercise_data.csv'):
    df = pd.read_csv('raw_exercise_data.csv')
    df['VO2_Diff'] = df['Post_VO2'] - df['Pre_VO2']
    df['Lactate_Diff'] = df['Post_Lactate'] - df['Pre_Lactate']
    summary = df.groupby('Group')[['Pre_VO2', 'Post_VO2', 'VO2_Diff', 'Pre_Lactate', 'Post_Lactate', 'Lactate_Diff']].agg(['mean', 'std']).round(2)
    
    # 표 생성
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Group'
    hdr_cells[1].text = 'Pre VO2max'
    hdr_cells[2].text = 'Post VO2max'
    hdr_cells[3].text = 'Pre Lactate'
    hdr_cells[4].text = 'Post Lactate'
    
    for group in ['Control', 'Training']:
        row_cells = table.add_row().cells
        row_cells[0].text = group
        row_cells[1].text = f"{summary.loc[group, ('Pre_VO2', 'mean')]} ± {summary.loc[group, ('Pre_VO2', 'std')]}"
        row_cells[2].text = f"{summary.loc[group, ('Post_VO2', 'mean')]} ± {summary.loc[group, ('Post_VO2', 'std')]}"
        row_cells[3].text = f"{summary.loc[group, ('Pre_Lactate', 'mean')]} ± {summary.loc[group, ('Pre_Lactate', 'std')]}"
        row_cells[4].text = f"{summary.loc[group, ('Post_Lactate', 'mean')]} ± {summary.loc[group, ('Post_Lactate', 'std')]}"

doc.add_paragraph().paragraph_format.space_after = Pt(12)

# 3. 통계 및 시각화 결과 삽입
doc.add_heading('2. 통계 검정 및 논문 시각화 결과', level=1)

figures = [
    ('lactate_threshold_figure.png', 'Figure 1. 점증적 운동검사 기반 젖산역치(Lactate Threshold) 및 VO2 응답 곡선'),
    ('stat_results_figure.png', 'Figure 2. 8주 운동 트레이닝 처치 전후 피험자 개별 VO2max 변화 추이 (Paired t-test, p < 0.001)'),
    ('anova_interaction_figure.png', 'Figure 3. 통제군과 운동군의 시기별 VO2max 상호작용 효과 (2-Way RM-ANOVA Interaction)')
]

for fig_path, caption in figures:
    if os.path.exists(fig_path):
        doc.add_paragraph(caption, style='List Bullet')
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(fig_path, width=Inches(5.5))
        p.paragraph_format.space_after = Pt(12)

# 문서 저장
output_docx = 'exercise_physiology_report.docx'
doc.save(output_docx)
print(f'[성공] MS Word 논문 보고서 빌드 완료: {output_docx}')
